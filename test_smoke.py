"""
Teste de fumaça: valida toda a aplicação (banco, exportações, API, fluxo de job)
SEM tocar no Google Maps de verdade — a função scrape_google_maps é substituída
por uma versão falsa que gera leads fictícios, já que este ambiente de teste
não tem acesso à internet pública.
"""
import asyncio
import os
import sys
import traceback

sys.path.insert(0, ".")

results = []


def check(name, fn):
    try:
        fn()
        results.append((name, True, ""))
        print(f"OK   - {name}")
    except Exception as e:
        results.append((name, False, str(e)))
        print(f"FAIL - {name}: {e}")
        traceback.print_exc()


# 1) Imports
def test_imports():
    import backend.config  # noqa
    import backend.database  # noqa
    import backend.models  # noqa
    import backend.scraper  # noqa
    import backend.jobs  # noqa
    import backend.webhook  # noqa
    import backend.exporters  # noqa
    import backend.main  # noqa


check("imports de todos os módulos", test_imports)

# 2) DB init + CRUD
import json
from datetime import datetime, timedelta, timezone
from backend.database import init_db, engine
from backend.models import Job, Lead, Schedule
from sqlmodel import Session, select

init_db()


def test_db_crud():
    with Session(engine) as session:
        job = Job(query="barbearia", location="São José do Rio Preto, SP", max_results=10, status="completed")
        session.add(job)
        session.commit()
        session.refresh(job)

        leads = [
            Lead(job_id=job.id, name="Barbearia do Zé", category="Barbearia",
                 address="Rua A, 123", phone="(17) 99999-0001", website=None,
                 rating=4.7, review_count=120, maps_url="https://maps.google.com/1"),
            Lead(job_id=job.id, name="Salão Bela Flor", category="Salão de beleza",
                 address="Rua B, 456", phone="(17) 99999-0002", website="https://belaflor.com.br",
                 rating=4.2, review_count=45, maps_url="https://maps.google.com/2"),
            Lead(job_id=job.id, name="Barber Shop Premium", category="Barbearia",
                 address="Av. C, 789", phone=None, website=None,
                 rating=5.0, review_count=8, maps_url="https://maps.google.com/3"),
        ]
        for l in leads:
            session.add(l)
        session.commit()

        found = session.exec(select(Lead).where(Lead.job_id == job.id)).all()
        assert len(found) == 3, f"esperava 3 leads, achei {len(found)}"
        global TEST_JOB_ID
        TEST_JOB_ID = job.id


check("banco de dados: criação de tabelas + CRUD", test_db_crud)

# 3) Exporters
from backend import exporters


def test_csv_export():
    with Session(engine) as session:
        csv_content = exporters.export_csv(session, TEST_JOB_ID)
        assert "Barbearia do Zé" in csv_content
        assert "Nome,Categoria" in csv_content
        lines = [l for l in csv_content.strip().split("\n") if l]
        assert len(lines) == 4, f"esperava 4 linhas (header+3), achei {len(lines)}"


check("exportação CSV", test_csv_export)


def test_xlsx_export():
    with Session(engine) as session:
        xlsx_bytes = exporters.export_xlsx(session, TEST_JOB_ID)
        assert xlsx_bytes[:2] == b"PK", "arquivo xlsx deveria começar com assinatura ZIP (PK)"
        assert len(xlsx_bytes) > 1000

        from openpyxl import load_workbook
        import io
        wb = load_workbook(io.BytesIO(xlsx_bytes))
        ws = wb.active
        assert ws.cell(1, 1).value == "Nome"
        assert ws.max_row == 4


check("exportação Excel (.xlsx)", test_xlsx_export)


def test_filter_no_website():
    with Session(engine) as session:
        pairs = exporters._get_leads(session, TEST_JOB_ID, only_no_website=True)
        assert len(pairs) == 2, f"esperava 2 leads sem site, achei {len(pairs)}"
        names = {lead.name for lead, _count in pairs}
        assert "Salão Bela Flor" not in names


check("filtro 'somente sem site'", test_filter_no_website)

# 4) FastAPI app + endpoints (com scraper mockado)
from fastapi.testclient import TestClient
import backend.jobs as jobs_module


async def _fake_scrape(query, location, max_results, log, on_lead, on_links_collected, on_link_result, should_cancel):
    await log(f"[fake] buscando {query} em {location}")
    hrefs = [f"https://maps.google.com/fake{i}" for i in range(3)]
    links_with_ids = await on_links_collected(hrefs)
    for i, (link_id, href) in enumerate(links_with_ids):
        await on_lead({
            "name": f"Estabelecimento Fake {i+1}",
            "category": "Teste",
            "address": "Rua Fake, 1",
            "phone": "(17) 90000-000" + str(i),
            "website": None if i % 2 == 0 else "https://exemplo.com",
            "rating": 4.5,
            "review_count": 10 * (i + 1),
            "maps_url": href,
            "latitude": -20.8,
            "longitude": -49.3,
        })
        await on_link_result(link_id, True)
    await log("[fake] concluído")
    return "completed"


def test_api_flow():
    jobs_module.scraper_module.scrape_google_maps = _fake_scrape
    import backend.main as main_module

    client = TestClient(main_module.app)

    r = client.get("/api/status")
    assert r.status_code == 200, r.text
    assert "n8n_webhook_configured" in r.json()

    r = client.post("/api/jobs", json={"query": "pizzaria", "location": "Rio Preto", "max_results": 5})
    assert r.status_code == 200, r.text
    job_id = r.json()["job_id"]
    assert job_id

    import time
    for _ in range(20):
        r = client.get(f"/api/jobs/{job_id}")
        if r.json()["status"] in ("completed", "failed"):
            break
        time.sleep(0.2)

    job_data = r.json()
    assert job_data["status"] == "completed", f"job terminou como: {job_data}"
    assert job_data["total_found"] == 3, job_data

    r = client.get(f"/api/leads?job_id={job_id}")
    assert r.status_code == 200
    assert len(r.json()) == 3

    r = client.get(f"/api/leads?job_id={job_id}&only_no_website=true")
    assert len(r.json()) == 2

    r = client.get(f"/api/export/csv?job_id={job_id}")
    assert r.status_code == 200
    assert "Estabelecimento Fake" in r.text

    r = client.get(f"/api/export/xlsx?job_id={job_id}")
    assert r.status_code == 200
    assert r.content[:2] == b"PK"

    r = client.get("/")
    assert r.status_code == 200
    assert "<html" in r.text.lower()


check("fluxo completo via API (job fake -> banco -> export)", test_api_flow)

# 5) API key protection
def test_api_key_protection():
    from backend.config import settings
    original = settings.API_KEY
    try:
        settings.API_KEY = "segredo123"
        import backend.main as main_module
        client = TestClient(main_module.app)

        r = client.get("/api/leads")
        assert r.status_code == 401, f"esperava 401 sem api key, veio {r.status_code}"

        r = client.get("/api/leads", headers={"X-API-Key": "segredo123"})
        assert r.status_code == 200, f"esperava 200 com api key certa, veio {r.status_code}"

        r = client.get("/api/leads", headers={"X-API-Key": "errada"})
        assert r.status_code == 401
    finally:
        settings.API_KEY = original


check("proteção por API key", test_api_key_protection)


# 5b) Robustez da raspagem: _is_blocked detecta URL/iframe/título sem falso positivo
class _FakeLocator:
    def __init__(self, count_val):
        self._count = count_val

    async def count(self):
        return self._count


class _FakePage:
    def __init__(self, url="https://maps.google.com/place/barbearia-xyz", title_text="", iframe_count=0):
        self.url = url
        self._title = title_text
        self._iframe_count = iframe_count

    async def title(self):
        return self._title

    def locator(self, selector):
        return _FakeLocator(self._iframe_count)


async def _test_is_blocked_async():
    from backend.scraper import _is_blocked

    assert await _is_blocked(_FakePage()) is False
    assert await _is_blocked(_FakePage(url="https://www.google.com/sorry/index?continue=x")) is True
    assert await _is_blocked(_FakePage(iframe_count=1)) is True
    assert await _is_blocked(_FakePage(title_text="Tráfego incomum detectado")) is True


def test_is_blocked_detection():
    asyncio.run(_test_is_blocked_async())


check("_is_blocked: detecta bloqueio por URL/iframe/título sem falso positivo", test_is_blocked_detection)


# 5c) Retry: recupera depois de falhas dentro do limite
async def _test_retry_succeeds_async():
    from backend.scraper import _extract_with_retry
    import backend.scraper as scraper_module

    original_delay = scraper_module._rand_delay
    scraper_module._rand_delay = lambda *a, **k: 0.001
    try:
        calls = {"n": 0}

        async def flaky(page, href):
            calls["n"] += 1
            if calls["n"] <= 2:
                raise Exception(f"falha simulada {calls['n']}")
            return {"name": "Sucesso", "maps_url": href}

        async def noop_log(msg):
            pass

        data = await _extract_with_retry(None, "http://x", noop_log, 1, 1, max_retries=3, extract_fn=flaky)
        assert data["name"] == "Sucesso"
        assert calls["n"] == 3
    finally:
        scraper_module._rand_delay = original_delay


def test_retry_succeeds_after_failures():
    asyncio.run(_test_retry_succeeds_async())


check("retry: recupera depois de falhas dentro do limite", test_retry_succeeds_after_failures)


# 5d) Retry: desiste com ExtractionFailed após esgotar tentativas
async def _test_retry_gives_up_async():
    from backend.scraper import _extract_with_retry, ExtractionFailed
    import backend.scraper as scraper_module

    original_delay = scraper_module._rand_delay
    scraper_module._rand_delay = lambda *a, **k: 0.001
    try:
        async def always_fail(page, href):
            raise Exception("sempre falha")

        async def noop_log(msg):
            pass

        try:
            await _extract_with_retry(None, "http://x", noop_log, 1, 1, max_retries=2, extract_fn=always_fail)
            assert False, "deveria ter levantado ExtractionFailed"
        except ExtractionFailed:
            pass
    finally:
        scraper_module._rand_delay = original_delay


def test_retry_gives_up_after_max_attempts():
    asyncio.run(_test_retry_gives_up_async())


check("retry: desiste com ExtractionFailed após esgotar tentativas", test_retry_gives_up_after_max_attempts)


# 5e) BlockedError propaga na hora, sem consumir tentativas de retry
async def _test_blocked_no_retry_async():
    from backend.scraper import _extract_with_retry, BlockedError
    import backend.scraper as scraper_module

    original_delay = scraper_module._rand_delay
    scraper_module._rand_delay = lambda *a, **k: 0.001
    try:
        calls = {"n": 0}

        async def blocked_fn(page, href):
            calls["n"] += 1
            raise BlockedError("bloqueado")

        async def noop_log(msg):
            pass

        try:
            await _extract_with_retry(None, "http://x", noop_log, 1, 1, max_retries=3, extract_fn=blocked_fn)
            assert False, "deveria ter levantado BlockedError"
        except BlockedError:
            pass
        assert calls["n"] == 1, f"BlockedError não deveria ter retry, mas foi chamado {calls['n']} vez(es)"
    finally:
        scraper_module._rand_delay = original_delay


def test_blocked_error_skips_retry():
    asyncio.run(_test_blocked_no_retry_async())


check("retry: BlockedError propaga na hora, sem retry", test_blocked_error_skips_retry)


# 5f) Circuit breaker: pausa ('partial') após N falhas seguidas
async def _test_circuit_breaker_async():
    from backend.scraper import _process_links
    import backend.scraper as scraper_module
    from backend.config import settings

    original_delay = scraper_module._rand_delay
    original_threshold = settings.CIRCUIT_BREAKER_THRESHOLD
    original_retries = settings.MAX_RETRIES_PER_LISTING
    scraper_module._rand_delay = lambda *a, **k: 0.001
    settings.CIRCUIT_BREAKER_THRESHOLD = 3
    settings.MAX_RETRIES_PER_LISTING = 1
    try:
        async def always_fail(page, href):
            raise Exception("falha simulada")

        async def noop_log(msg):
            pass

        leads = []

        async def on_lead(data):
            leads.append(data)

        results = []

        async def on_link_result(link_id, success):
            results.append((link_id, success))

        def should_cancel():
            return False

        links_with_ids = [(i, f"http://x{i}") for i in range(1, 11)]
        status = await _process_links(
            None, links_with_ids, noop_log, on_lead, on_link_result, should_cancel,
            extract_fn=always_fail,
        )
        assert status == "partial", f"esperava 'partial', veio '{status}'"
        assert len(results) == 3, f"esperava parar em 3 links processados, processou {len(results)}"
        assert all(success is False for _, success in results)
        assert len(leads) == 0
    finally:
        scraper_module._rand_delay = original_delay
        settings.CIRCUIT_BREAKER_THRESHOLD = original_threshold
        settings.MAX_RETRIES_PER_LISTING = original_retries


def test_circuit_breaker_pauses_after_threshold():
    asyncio.run(_test_circuit_breaker_async())


check("circuit breaker: pausa a varredura ('partial') após N falhas seguidas", test_circuit_breaker_pauses_after_threshold)


# 5g) Bloqueio detectado interrompe a varredura inteira ('blocked')
async def _test_blocked_stops_all_async():
    from backend.scraper import _process_links, BlockedError
    import backend.scraper as scraper_module

    original_delay = scraper_module._rand_delay
    scraper_module._rand_delay = lambda *a, **k: 0.001
    try:
        calls = {"n": 0}

        async def blocked_after_two(page, href):
            calls["n"] += 1
            if calls["n"] <= 2:
                return {"name": f"Lead {calls['n']}", "maps_url": href}
            raise BlockedError("bloqueado")

        leads = []

        async def on_lead(data):
            leads.append(data)

        results = []

        async def on_link_result(link_id, success):
            results.append((link_id, success))

        async def noop_log(msg):
            pass

        def should_cancel():
            return False

        links_with_ids = [(i, f"http://x{i}") for i in range(1, 6)]
        status = await _process_links(
            None, links_with_ids, noop_log, on_lead, on_link_result, should_cancel,
            extract_fn=blocked_after_two,
        )
        assert status == "blocked"
        assert len(leads) == 2
        assert len(results) == 3
        assert results[-1] == (3, False)
    finally:
        scraper_module._rand_delay = original_delay


def test_blocked_stops_entire_job():
    asyncio.run(_test_blocked_stops_all_async())


check("bloqueio detectado interrompe a varredura inteira ('blocked')", test_blocked_stops_entire_job)


# 5h) Cancelamento no meio da varredura para com status 'cancelled'
async def _test_cancel_mid_loop_async():
    from backend.scraper import _process_links

    async def ok_fn(page, href):
        return {"name": "Lead", "maps_url": href}

    results = []

    async def on_link_result(link_id, success):
        results.append((link_id, success))

    state = {"cancelled": False}

    def should_cancel():
        return state["cancelled"]

    leads = []

    async def on_lead_and_cancel(data):
        leads.append(data)
        if len(leads) == 2:
            state["cancelled"] = True

    async def noop_log(msg):
        pass

    links_with_ids = [(i, f"http://x{i}") for i in range(1, 6)]
    status = await _process_links(
        None, links_with_ids, noop_log, on_lead_and_cancel, on_link_result, should_cancel,
        extract_fn=ok_fn,
    )
    assert status == "cancelled"
    assert len(leads) == 2


def test_cancellation_stops_processing():
    asyncio.run(_test_cancel_mid_loop_async())


check("cancelamento no meio da varredura para com status 'cancelled'", test_cancellation_stops_processing)


# 5i) Varredura sem falhas termina 'completed'
async def _test_clean_completion_async():
    from backend.scraper import _process_links

    async def ok_fn(page, href):
        return {"name": f"Lead {href}", "maps_url": href}

    leads = []

    async def on_lead(data):
        leads.append(data)

    results = []

    async def on_link_result(link_id, success):
        results.append((link_id, success))

    async def noop_log(msg):
        pass

    def should_cancel():
        return False

    links_with_ids = [(i, f"http://x{i}") for i in range(1, 4)]
    status = await _process_links(
        None, links_with_ids, noop_log, on_lead, on_link_result, should_cancel, extract_fn=ok_fn,
    )
    assert status == "completed"
    assert len(leads) == 3
    assert all(success for _, success in results)


def test_clean_run_completes():
    asyncio.run(_test_clean_completion_async())


check("varredura sem falhas termina com status 'completed'", test_clean_run_completes)


# 5j) Fluxo de resume ponta a ponta (via jobs.py, com scraper falso)
async def _test_resume_flow_async():
    import backend.jobs as jobs_module
    from backend.models import JobLink

    with Session(engine) as session:
        job = Job(query="teste-resume", location="teste", max_results=10, status="pending")
        session.add(job)
        session.commit()
        session.refresh(job)
        job_id = job.id

    original_scrape = jobs_module.scraper_module.scrape_google_maps
    original_resume = jobs_module.scraper_module.resume_scrape

    async def fake_scrape_partial(query, location, max_results, log, on_lead, on_links_collected, on_link_result, should_cancel):
        await log("fake: coletando links")
        links_with_ids = await on_links_collected(["http://a", "http://b", "http://c", "http://d"])
        await on_lead({"name": "Lead A", "maps_url": "http://a"})
        await on_link_result(links_with_ids[0][0], True)
        await on_lead({"name": "Lead B", "maps_url": "http://b"})
        await on_link_result(links_with_ids[1][0], True)
        # links_with_ids[2] e [3] permanecem "pending" de propósito — é assim que
        # um circuit breaker de verdade deixaria os itens que nunca chegou a tentar
        return "partial"

    async def fake_resume_completes(links_with_ids, log, on_lead, on_link_result, should_cancel):
        await log("fake: retomando")
        for link_id, href in links_with_ids:
            await on_lead({"name": f"Lead recuperado {href}", "maps_url": href})
            await on_link_result(link_id, True)
        return "completed"

    jobs_module.scraper_module.scrape_google_maps = fake_scrape_partial
    jobs_module.scraper_module.resume_scrape = fake_resume_completes

    try:
        await jobs_module.run_job(job_id, "teste-resume", "teste", 10)

        with Session(engine) as session:
            job = session.get(Job, job_id)
            assert job.status == "partial", f"esperava 'partial', veio '{job.status}'"
            assert job.total_found == 2

        assert jobs_module.is_resumable(job_id) is True
        assert jobs_module.pending_count(job_id) == 2

        await jobs_module.resume_job(job_id)

        with Session(engine) as session:
            job = session.get(Job, job_id)
            assert job.status == "completed", f"esperava 'completed', veio '{job.status}'"
            assert job.total_found == 4

        assert jobs_module.is_resumable(job_id) is False
    finally:
        jobs_module.scraper_module.scrape_google_maps = original_scrape
        jobs_module.scraper_module.resume_scrape = original_resume
        with Session(engine) as session:
            for l in session.exec(select(Lead).where(Lead.job_id == job_id)).all():
                session.delete(l)
            for jl in session.exec(select(JobLink).where(JobLink.job_id == job_id)).all():
                session.delete(jl)
            j = session.get(Job, job_id)
            if j:
                session.delete(j)
            session.commit()


def test_resume_flow_end_to_end():
    asyncio.run(_test_resume_flow_async())


check("fluxo de resume: job 'partial' -> retomado -> 'completed'", test_resume_flow_end_to_end)


# 5k) API: /resume e /resumable ponta a ponta
def test_api_resume_flow():
    import backend.jobs as jobs_module
    import backend.main as main_module

    async def fake_scrape_partial(query, location, max_results, log, on_lead, on_links_collected, on_link_result, should_cancel):
        links_with_ids = await on_links_collected(["http://p1", "http://p2", "http://p3"])
        await on_lead({"name": "Lead 1", "maps_url": "http://p1"})
        await on_link_result(links_with_ids[0][0], True)
        # links_with_ids[1] e [2] permanecem "pending" (circuito nunca os alcançou)
        return "partial"

    async def fake_resume_completes(links_with_ids, log, on_lead, on_link_result, should_cancel):
        for link_id, href in links_with_ids:
            await on_lead({"name": "Lead recuperado", "maps_url": href})
            await on_link_result(link_id, True)
        return "completed"

    original_scrape = jobs_module.scraper_module.scrape_google_maps
    original_resume = jobs_module.scraper_module.resume_scrape
    jobs_module.scraper_module.scrape_google_maps = fake_scrape_partial
    jobs_module.scraper_module.resume_scrape = fake_resume_completes

    try:
        client = TestClient(main_module.app)

        r = client.post("/api/jobs", json={"query": "x", "location": "y", "max_results": 5})
        assert r.status_code == 200, r.text
        job_id = r.json()["job_id"]

        import time
        for _ in range(20):
            r = client.get(f"/api/jobs/{job_id}")
            if r.json()["status"] in ("completed", "failed", "partial", "blocked", "cancelled"):
                break
            time.sleep(0.2)
        assert r.json()["status"] == "partial", r.json()

        r = client.get(f"/api/jobs/{job_id}/resumable")
        assert r.status_code == 200
        assert r.json()["resumable"] is True
        assert r.json()["pending_count"] == 2

        r = client.post(f"/api/jobs/{job_id}/resume")
        assert r.status_code == 200, r.text

        for _ in range(20):
            r = client.get(f"/api/jobs/{job_id}")
            if r.json()["status"] in ("completed", "failed"):
                break
            time.sleep(0.2)
        assert r.json()["status"] == "completed", r.json()
        assert r.json()["total_found"] == 3

        r = client.get(f"/api/jobs/{job_id}/resumable")
        assert r.json()["resumable"] is False

        r = client.post(f"/api/jobs/{job_id}/resume")
        assert r.status_code == 400
    finally:
        jobs_module.scraper_module.scrape_google_maps = original_scrape
        jobs_module.scraper_module.resume_scrape = original_resume


check("API: POST /resume e GET /resumable funcionam ponta a ponta", test_api_resume_flow)


# 6) Regressão: reset_queue esvazia a fila em vez de trocar o objeto (bug da
# corrida no /stream — ver comentário em jobs.py:reset_queue)
def test_reset_queue_keeps_same_object_and_drains():
    import asyncio as _asyncio
    from backend.jobs import get_queue, reset_queue, _queues

    async def _run():
        job_id = "teste-reset-queue"
        _queues.pop(job_id, None)
        q1 = get_queue(job_id)
        await q1.put({"type": "log", "message": "mensagem antiga, deveria sumir"})
        q2 = reset_queue(job_id)
        assert q1 is q2, "reset_queue não deveria trocar o objeto Queue"
        assert q2.empty(), "reset_queue deveria ter esvaziado a fila"

    _asyncio.run(_run())


check("reset_queue: mantém o mesmo objeto e esvazia mensagens antigas", test_reset_queue_keeps_same_object_and_drains)


# 7) Regressão: job 'running' órfão (servidor derrubado no meio) vira 'partial' na subida
def test_recover_orphaned_jobs():
    from backend.jobs import recover_orphaned_jobs

    with Session(engine) as session:
        job = Job(query="orfao", location="teste", status="running")
        session.add(job)
        session.commit()
        session.refresh(job)
        job_id = job.id

    n = recover_orphaned_jobs()
    assert n >= 1, f"esperava recuperar ao menos 1 job órfão, recuperou {n}"

    with Session(engine) as session:
        job = session.get(Job, job_id)
        assert job.status == "partial", f"esperava 'partial', veio '{job.status}'"
        assert job.finished_at is not None


check("recover_orphaned_jobs: job preso em 'running' vira 'partial' na subida", test_recover_orphaned_jobs)


# 8) Regressão: não deixa criar um job novo enquanto outro está 'running'
def test_blocks_concurrent_jobs():
    import backend.main as main_module
    client = TestClient(main_module.app)

    with Session(engine) as session:
        job = Job(query="em-andamento", location="teste", status="running")
        session.add(job)
        session.commit()
        session.refresh(job)
        job_id = job.id

    try:
        r = client.post("/api/jobs", json={"query": "outra", "location": "cidade", "max_results": 5})
        assert r.status_code == 409, f"esperava 409 com job já rodando, veio {r.status_code}: {r.text}"
    finally:
        with Session(engine) as session:
            j = session.get(Job, job_id)
            if j:
                session.delete(j)
                session.commit()


check("POST /api/jobs: retorna 409 se já existe varredura em andamento", test_blocks_concurrent_jobs)


# 9) Regressão: exportação CSV começa com BOM UTF-8 (acentuação correta no Excel)
def test_csv_export_has_bom():
    with Session(engine) as session:
        csv_content = exporters.export_csv(session, TEST_JOB_ID)
        assert csv_content.startswith("\ufeff"), "CSV deveria começar com BOM UTF-8"
        assert csv_content[1:].startswith("Nome,Categoria"), "cabeçalho deveria vir logo após o BOM"


check("exportação CSV: começa com BOM UTF-8", test_csv_export_has_bom)


# 10) Regressão: nome de arquivo de exportação é dinâmico (nicho + local + data)
def test_export_filename_is_dynamic():
    with Session(engine) as session:
        name = exporters.export_filename(session, TEST_JOB_ID, "csv")
        assert name.endswith(".csv")
        assert name != "leads.csv", "nome do arquivo deveria refletir a busca, não ser genérico"
        assert "barbearia" in name


check("export_filename: gera nome dinâmico a partir da busca", test_export_filename_is_dynamic)


# 11) Regressão: evento SSE "lead" carrega o lead completo (id, job_id, scraped_at),
# não só o dict cru do scraper — necessário pra poder excluir um lead ao vivo
def test_sse_lead_event_has_full_payload():
    import backend.jobs as jobs_module

    async def _run():
        job_id = "teste-sse-payload"
        with Session(engine) as session:
            session.add(Job(id=job_id, query="x", location="y", status="running"))
            session.commit()

        log, on_lead, _links, _result, _cancel = jobs_module._make_callbacks(job_id)
        await on_lead({"name": "Estabelecimento X", "maps_url": "http://x"})

        q = jobs_module.get_queue(job_id)
        msg = await asyncio.wait_for(q.get(), timeout=2.0)
        assert msg["type"] == "lead"
        lead = msg["lead"]
        assert lead.get("id") is not None, "evento SSE 'lead' deveria incluir o id do banco"
        assert lead.get("job_id") == job_id
        assert lead.get("scraped_at") is not None
        assert lead.get("name") == "Estabelecimento X"

    asyncio.run(_run())
    with Session(engine) as session:
        for l in session.exec(select(Lead).where(Lead.job_id == "teste-sse-payload")).all():
            session.delete(l)
        j = session.get(Job, "teste-sse-payload")
        if j:
            session.delete(j)
        session.commit()


check("evento SSE 'lead': payload completo com id/job_id/scraped_at", test_sse_lead_event_has_full_payload)

# 12) Regressão: variável presente-porém-vazia no .env não pode virar string vazia
def test_empty_env_var_falls_back_to_default():
    from backend.config import _get, _bool

    os.environ["TESTE_VAZIO"] = ""
    os.environ["TESTE_PREENCHIDO"] = "valor-real"
    os.environ.pop("TESTE_AUSENTE", None)
    try:
        assert _get("TESTE_VAZIO", "padrao") == "padrao"
        assert _get("TESTE_PREENCHIDO", "padrao") == "valor-real"
        assert _get("TESTE_AUSENTE", "padrao") == "padrao"
        assert _bool("TESTE_VAZIO", True) is True
    finally:
        os.environ.pop("TESTE_VAZIO", None)
        os.environ.pop("TESTE_PREENCHIDO", None)


check("regressão: variável vazia no .env cai no padrão (_get/_bool)", test_empty_env_var_falls_back_to_default)


# 12b) Regressão: datas voltam do banco com tzinfo=UTC (não "ingênuas") e viram
# JSON com o fuso explícito — sem isso, comparações em Python quebram com
# TypeError e o JavaScript do navegador interpreta a data como hora LOCAL em
# vez de UTC, mostrando horários errados na tela.
def test_datetime_roundtrip_keeps_utc_tzinfo():
    with Session(engine) as session:
        job = Job(query="tz-teste", location="x")
        session.add(job)
        session.commit()
        job_id = job.id

    try:
        with Session(engine) as session:  # sessão NOVA, força reler do banco
            fetched = session.get(Job, job_id)
            assert fetched.created_at.tzinfo is not None, "data voltou 'ingênua' (sem tzinfo) do banco"
            assert fetched.created_at.tzinfo == timezone.utc

            # não pode levantar TypeError ("can't compare offset-naive and offset-aware")
            assert fetched.created_at <= datetime.now(timezone.utc)

            payload = fetched.model_dump(mode="json")
            assert payload["created_at"].endswith("Z") or "+00:00" in payload["created_at"], (
                f"JSON sem fuso explícito, JS vai interpretar como hora local: {payload['created_at']!r}"
            )
    finally:
        with Session(engine) as session:
            j = session.get(Job, job_id)
            if j:
                session.delete(j)
                session.commit()


check("datas: leitura do banco mantém tzinfo=UTC e serializa com fuso explícito no JSON", test_datetime_roundtrip_keeps_utc_tzinfo)


# 13) Regressão: migração leve adiciona a coluna 'status' num banco 'antigo'
# (create_all não altera tabelas já existentes, só cria as que faltam)
def test_schema_migration_adds_status_column():
    import sqlite3
    import tempfile
    from backend.database import _ensure_schema_upgrades

    fd, path = tempfile.mkstemp(suffix=".db")
    os.close(fd)
    try:
        conn = sqlite3.connect(path)
        conn.execute(
            "CREATE TABLE lead (id INTEGER PRIMARY KEY, job_id TEXT, name TEXT, "
            "maps_url TEXT, scraped_at TEXT)"
        )
        conn.execute("INSERT INTO lead (id, job_id, name, maps_url, scraped_at) VALUES (1,'j','N','u','t')")
        conn.commit()
        conn.close()

        import backend.database as db_module
        original_engine = db_module.engine
        from sqlmodel import create_engine
        db_module.engine = create_engine(f"sqlite:///{path}", connect_args={"check_same_thread": False})
        try:
            _ensure_schema_upgrades()
            conn = sqlite3.connect(path)
            cols = [r[1] for r in conn.execute("PRAGMA table_info(lead)").fetchall()]
            assert "status" in cols, "migração deveria ter adicionado a coluna 'status'"
            row = conn.execute("SELECT status FROM lead WHERE id=1").fetchone()
            assert row[0] == "novo", f"esperava default 'novo', veio {row[0]!r}"
            conn.close()
        finally:
            db_module.engine = original_engine
    finally:
        os.remove(path)


check("migração de schema: banco antigo ganha a coluna 'status' do Lead", test_schema_migration_adds_status_column)


# 14) Regressão: CRM leve — atualizar e validar o status de um lead
def test_lead_status_update():
    import backend.main as main_module
    client = TestClient(main_module.app)

    with Session(engine) as session:
        lead = Lead(job_id=TEST_JOB_ID, name="Lead pra status", maps_url="http://status-test")
        session.add(lead)
        session.commit()
        session.refresh(lead)
        lead_id = lead.id

    try:
        r = client.patch(f"/api/leads/{lead_id}/status", json={"status": "contatado"})
        assert r.status_code == 200, r.text
        assert r.json()["status"] == "contatado"

        r = client.patch(f"/api/leads/{lead_id}/status", json={"status": "status-que-nao-existe"})
        assert r.status_code == 400, "status inválido deveria ser rejeitado"

        r = client.patch("/api/leads/99999999/status", json={"status": "fechado"})
        assert r.status_code == 404
    finally:
        with Session(engine) as session:
            lead = session.get(Lead, lead_id)
            if lead:
                session.delete(lead)
                session.commit()


check("CRM leve: PATCH /api/leads/{id}/status valida e atualiza", test_lead_status_update)


# 15) Regressão: deduplicação de leads entre buscas diferentes (via API)
def test_dedupe_across_jobs_api():
    import backend.main as main_module
    client = TestClient(main_module.app)

    with Session(engine) as session:
        job_a = Job(id="dedupA", query="x", location="y", status="completed")
        job_b = Job(id="dedupB", query="x", location="y", status="completed")
        session.add(job_a)
        session.add(job_b)
        same_url = "https://maps.google.com/place/@0,0,0z/data=!1s0xdead:0xbeef"
        session.add(Lead(job_id="dedupA", name="Achado Duas Vezes", address="Rua X",
                          maps_url=same_url + "?a=1"))
        session.add(Lead(job_id="dedupB", name="Achado Duas Vezes", address="Rua X",
                          maps_url=same_url + "?a=2", website="https://site.com"))
        session.add(Lead(job_id="dedupA", name="Único", address="Rua Y", maps_url="http://unico"))
        session.commit()

    try:
        r = client.get("/api/leads?dedupe=true")
        assert r.status_code == 200, r.text
        data = r.json()
        by_name = {d["name"]: d for d in data if d["name"] in ("Achado Duas Vezes", "Único")}
        assert by_name["Achado Duas Vezes"]["duplicate_count"] == 2
        assert by_name["Achado Duas Vezes"]["website"] == "https://site.com", "deveria escolher o representante mais completo"
        assert by_name["Único"]["duplicate_count"] == 1

        r_normal = client.get("/api/leads?job_id=dedupA")
        assert "duplicate_count" not in r_normal.json()[0], "sem dedupe=true, não deveria ter duplicate_count"
    finally:
        with Session(engine) as session:
            for l in session.exec(select(Lead).where(Lead.job_id.in_(["dedupA", "dedupB"]))).all():
                session.delete(l)
            for j in (session.get(Job, "dedupA"), session.get(Job, "dedupB")):
                if j:
                    session.delete(j)
            session.commit()


check("dedupe entre buscas: GET /api/leads?dedupe=true agrupa e escolhe o mais completo", test_dedupe_across_jobs_api)


# 16) Regressão: CRUD de agendamentos (buscas recorrentes)
def test_schedules_crud():
    import backend.main as main_module
    client = TestClient(main_module.app)

    r = client.post("/api/schedules", json={"query": "pizzaria", "location": "centro", "frequency_days": 7})
    assert r.status_code == 200, r.text
    sched = r.json()
    schedule_id = sched["id"]
    assert sched["active"] is True
    assert sched["frequency_days"] == 7

    try:
        r = client.get("/api/schedules")
        assert any(s["id"] == schedule_id for s in r.json())

        r = client.patch(f"/api/schedules/{schedule_id}", json={"active": False})
        assert r.status_code == 200
        assert r.json()["active"] is False

        r = client.patch(f"/api/schedules/{schedule_id}", json={"frequency_days": 14})
        assert r.json()["frequency_days"] == 14
    finally:
        r = client.delete(f"/api/schedules/{schedule_id}")
        assert r.status_code == 200
        r = client.get("/api/schedules")
        assert not any(s["id"] == schedule_id for s in r.json())


check("agendamentos: CRUD completo (criar, listar, pausar, mudar frequência, remover)", test_schedules_crud)


# 17) Regressão: o agendador dispara um job quando vencido, avança next_run_at,
# e NÃO dispara se já houver uma varredura em andamento
def test_scheduler_fires_when_due():
    async def _test():
        with Session(engine) as session:
            sched = Schedule(
                query="salão", location="centro", max_results=5, frequency_days=3,
                next_run_at=datetime.now(timezone.utc) - timedelta(minutes=1),  # já vencido
            )
            session.add(sched)
            session.commit()
            session.refresh(sched)
            sched_id = sched.id

        try:
            await jobs_module._run_due_schedules_once()
            await asyncio.sleep(0.05)  # dá tempo da task criada rodar até tentar abrir o job

            with Session(engine) as session:
                sched = session.get(Schedule, sched_id)
                assert sched.last_run_at is not None, "deveria ter marcado last_run_at"
                assert sched.last_job_id is not None, "deveria ter associado o job criado"
                assert sched.next_run_at > datetime.now(timezone.utc), "deveria ter avançado next_run_at"
                job = session.get(Job, sched.last_job_id)
                assert job is not None and job.query == "salão"

            # Não deve disparar de novo agora que o job criado está 'running'
            # (ou já mudou de status, dependendo da velocidade do teste) —
            # o importante aqui é que rodar de novo não quebra nem duplica.
            before = sched.last_run_at
            await jobs_module._run_due_schedules_once()
            with Session(engine) as session:
                sched = session.get(Schedule, sched_id)
                assert sched.last_run_at == before, "não deveria ter rodado de novo tão cedo"
        finally:
            with Session(engine) as session:
                sched = session.get(Schedule, sched_id)
                job_id = sched.last_job_id if sched else None
                if sched:
                    session.delete(sched)
                if job_id:
                    job = session.get(Job, job_id)
                    if job:
                        for l in session.exec(select(Lead).where(Lead.job_id == job_id)).all():
                            session.delete(l)
                        session.delete(job)
                session.commit()

    asyncio.run(_test())


check("agendador: dispara job vencido, avança next_run_at, evita concorrência", test_scheduler_fires_when_due)


# 18) Regressão: places_api.py — mapeamento de campos, paginação e erros,
# usando um transporte HTTP simulado (sem rede real nem chave de verdade)
def test_places_api_pagination_and_mapping():
    import httpx as httpx_module
    from backend import places_api
    from backend.config import settings as settings_module

    original_key = settings_module.GOOGLE_PLACES_API_KEY
    settings_module.GOOGLE_PLACES_API_KEY = "fake-key-de-teste"

    page1 = {
        "places": [
            {
                "id": "ChIJ_1", "displayName": {"text": "Barbearia do Zé"},
                "formattedAddress": "Rua A, 123", "nationalPhoneNumber": "(17) 99999-0001",
                "types": ["barber_shop"], "rating": 4.7, "userRatingCount": 120,
                "location": {"latitude": -20.81, "longitude": -49.38},
            },
            {"displayName": {"text": "Sem ID, deve ser pulado"}},
        ],
        "nextPageToken": "pagina2",
    }
    page2 = {"places": [{"id": "ChIJ_2", "displayName": {"text": "Studio Corte Fino"}}]}

    def handler(request):
        body = json.loads(request.content)
        assert request.headers["X-Goog-Api-Key"] == "fake-key-de-teste"
        if "pageToken" in body:
            return httpx_module.Response(200, json=page2)
        return httpx_module.Response(200, json=page1)

    original_client = httpx_module.AsyncClient

    def fake_client(*a, **kw):
        kw["transport"] = httpx_module.MockTransport(handler)
        return original_client(*a, **kw)

    async def _test():
        leads = []
        async def log(msg): pass
        async def on_lead(data): leads.append(data)
        async def on_links_collected(hrefs): return []
        async def on_link_result(lid, ok): pass
        def should_cancel(): return False

        places_api.httpx.AsyncClient = fake_client
        try:
            status = await places_api.search_google_places(
                "barbearia", "São José do Rio Preto, SP", 10,
                log, on_lead, on_links_collected, on_link_result, should_cancel,
            )
        finally:
            places_api.httpx.AsyncClient = original_client
        return status, leads

    try:
        status, leads = asyncio.run(_test())
        assert status == "completed"
        assert len(leads) == 2, f"esperava 2 leads (1 pulado por falta de id), veio {len(leads)}"
        assert leads[0]["name"] == "Barbearia do Zé"
        assert leads[0]["phone"] == "(17) 99999-0001"
        assert leads[0]["rating"] == 4.7
        assert leads[0]["maps_url"] == "https://www.google.com/maps/place/?q=place_id:ChIJ_1"
        assert leads[1]["name"] == "Studio Corte Fino"
    finally:
        settings_module.GOOGLE_PLACES_API_KEY = original_key


check("places_api: paginação, mapeamento de campos e item sem id é pulado", test_places_api_pagination_and_mapping)


def test_places_api_error_handling():
    import httpx as httpx_module
    from backend import places_api
    from backend.config import settings as settings_module

    original_key = settings_module.GOOGLE_PLACES_API_KEY
    settings_module.GOOGLE_PLACES_API_KEY = "fake-key"
    original_client = httpx_module.AsyncClient

    async def _run(handler, max_results=10):
        def fake_client(*a, **kw):
            kw["transport"] = httpx_module.MockTransport(handler)
            return original_client(*a, **kw)
        leads = []
        async def log(msg): pass
        async def on_lead(data): leads.append(data)
        async def on_links_collected(hrefs): return []
        async def on_link_result(lid, ok): pass
        def should_cancel(): return False
        places_api.httpx.AsyncClient = fake_client
        try:
            status = await places_api.search_google_places(
                "x", "y", max_results, log, on_lead, on_links_collected, on_link_result, should_cancel)
        finally:
            places_api.httpx.AsyncClient = original_client
        return status, leads

    try:
        # 403 -> failed, sem leads
        status, leads = asyncio.run(_run(lambda req: httpx_module.Response(403, json={"error": "nope"})))
        assert status == "failed" and len(leads) == 0

        # 429 depois de 1 página coletada -> partial, mantém o que já tinha
        call_n = {"n": 0}
        def handler_429(req):
            call_n["n"] += 1
            if call_n["n"] == 1:
                return httpx_module.Response(200, json={"places": [{"id": "a", "displayName": {"text": "A"}}], "nextPageToken": "x"})
            return httpx_module.Response(429, json={"error": "quota"})
        status, leads = asyncio.run(_run(handler_429, max_results=50))
        assert status == "partial" and len(leads) == 1

        # erro de rede -> failed
        def handler_net_error(req):
            raise httpx_module.ConnectError("sem rede", request=req)
        status, leads = asyncio.run(_run(handler_net_error))
        assert status == "failed" and len(leads) == 0
    finally:
        settings_module.GOOGLE_PLACES_API_KEY = original_key


check("places_api: 403/429/erro de rede tratados sem crashar", test_places_api_error_handling)


# 13) Regressão E2E: sobe o servidor uvicorn DE VERDADE (igual o run.bat faz),
# com .env copiado do .env.example, e faz um POST /api/jobs real por HTTP.
# Isso é o que de fato reproduz o bug original: rotas síncronas do FastAPI são
# despachadas pra threads de um threadpool, diferentes da thread principal
# que roda o init_db() no startup — um teste single-thread não pega isso.
def test_real_server_setup_env_example():
    import subprocess
    import shutil
    import time
    import urllib.request
    import urllib.error
    import json as jsonlib

    project_dir = os.path.abspath(".")
    env_path = os.path.join(project_dir, ".env")
    backup_path = env_path + ".bak"
    had_env = os.path.exists(env_path)
    if had_env:
        shutil.move(env_path, backup_path)

    db_path = os.path.join(project_dir, "data", "prospector.db")
    proc = None
    try:
        shutil.copy(os.path.join(project_dir, ".env.example"), env_path)
        if os.path.exists(db_path):
            os.remove(db_path)

        proc = subprocess.Popen(
            [os.path.join(project_dir, "venv", "bin", "python"), "-m", "uvicorn",
             "backend.main:app", "--host", "127.0.0.1", "--port", "8123"],
            cwd=project_dir, stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True,
        )

        up = False
        for _ in range(60):
            try:
                urllib.request.urlopen("http://127.0.0.1:8123/api/status", timeout=1)
                up = True
                break
            except Exception:
                time.sleep(0.25)
        if not up:
            out = proc.stdout.read() if proc.stdout else ""
            raise AssertionError(f"servidor não respondeu a tempo. saída:\n{out}")

        req = urllib.request.Request(
            "http://127.0.0.1:8123/api/jobs",
            data=jsonlib.dumps({"query": "barbearia", "location": "teste", "max_results": 10}).encode(),
            headers={"Content-Type": "application/json"},
            method="POST",
        )
        try:
            resp = urllib.request.urlopen(req, timeout=5)
            body = resp.read().decode()
            assert resp.status == 200, body
        except urllib.error.HTTPError as e:
            raise AssertionError(f"POST /api/jobs falhou: {e.code} {e.read().decode()}")

        job_id = jsonlib.loads(body)["job_id"]
        assert job_id, "resposta sem job_id"

        r2 = urllib.request.urlopen(f"http://127.0.0.1:8123/api/jobs/{job_id}", timeout=5)
        assert r2.status == 200
    finally:
        if proc:
            proc.terminate()
            try:
                proc.wait(timeout=5)
            except subprocess.TimeoutExpired:
                proc.kill()
        if os.path.exists(env_path):
            os.remove(env_path)
        if had_env:
            shutil.move(backup_path, env_path)
        if os.path.exists(db_path):
            os.remove(db_path)


check("regressão E2E (servidor real): POST /api/jobs com .env do exemplo", test_real_server_setup_env_example)



print("\n" + "=" * 50)
passed = sum(1 for _, ok, _ in results if ok)
print(f"{passed}/{len(results)} testes passaram")
if passed != len(results):
    print("\nFALHAS:")
    for name, ok, err in results:
        if not ok:
            print(f"  - {name}: {err}")
    sys.exit(1)
